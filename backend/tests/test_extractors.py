import io
import zipfile

import httpx
import pytest
import respx
from PIL import Image, ImageDraw, ImageFont

from app import extractors, ingest

needs_ocr = pytest.mark.skipif(not extractors.ocr_available(), reason="tesseract not installed")


def text_image(text: str) -> Image.Image:
    img = Image.new("RGB", (1400, 260), "white")
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 64)
    except OSError:
        font = ImageFont.load_default(size=64)
    draw.text((40, 90), text, fill="black", font=font)
    return img


def test_pptx_slides_and_notes():
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Mitochondria"
    slide.placeholders[1].text = "Make ATP by oxidative phosphorylation"
    slide.notes_slide.notes_text_frame.text = "Mention the proton gradient"
    prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = "Recap"
    buf = io.BytesIO()
    prs.save(buf)
    pages = extractors.extract_pptx(buf.getvalue())
    assert pages[0][0] == 1 and "Make ATP" in pages[0][1] and "Speaker notes: Mention the proton gradient" in pages[0][1]
    assert pages[1] == (2, "Recap")


def make_epub() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("mimetype", "application/epub+zip")
        z.writestr("META-INF/container.xml", '<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
        z.writestr("OEBPS/content.opf", '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>The Cell</dc:title></metadata><manifest><item id="c1" href="ch1.xhtml" media-type="application/xhtml+xml"/><item id="c2" href="text/ch2.xhtml" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="c2"/><itemref idref="c1"/></spine></package>')
        z.writestr("OEBPS/ch1.xhtml", "<html><body><h1>Chapter One</h1><p>Glycolysis happens in the cytoplasm.</p><script>bad()</script></body></html>")
        z.writestr("OEBPS/text/ch2.xhtml", "<html><body><h1>Preface</h1><p>Why cells need energy.</p></body></html>")
    return buf.getvalue()


def test_epub_follows_spine_order():
    title, chapters = extractors.extract_epub(make_epub())
    assert title == "The Cell"
    assert chapters[0] == (1, "Preface\n\nWhy cells need energy.")
    assert "Glycolysis happens in the cytoplasm." in chapters[1][1] and "bad()" not in chapters[1][1]


@pytest.mark.parametrize("url,expected", [
    ("https://docs.google.com/document/d/AbC_123/edit?usp=sharing", ("https://docs.google.com/document/d/AbC_123/export?format=txt", "text")),
    ("https://docs.google.com/presentation/d/XyZ/edit#slide=id.p", ("https://docs.google.com/presentation/d/XyZ/export/txt", "text")),
    ("https://docs.google.com/spreadsheets/d/S1/edit#gid=0", ("https://docs.google.com/spreadsheets/d/S1/export?format=csv", "text")),
    ("https://drive.google.com/file/d/F9/view", ("https://drive.google.com/uc?export=download&id=F9", "file")),
    ("https://example.com/document/d/x", None),
])
def test_google_export_urls(url, expected):
    assert extractors.google_export_url(url) == expected


async def test_fetch_google_doc_uses_export_and_filename():
    with respx.mock() as router:
        router.get("https://docs.google.com/document/d/D1/export?format=txt").mock(return_value=httpx.Response(
            200, text="Lecture notes on ATP synthase.", headers={"content-type": "text/plain; charset=utf-8",
                                                                  "content-disposition": "attachment; filename=\"Bio Notes.txt\""}))
        title, text, pages = await ingest.fetch_web("https://docs.google.com/document/d/D1/edit")
    assert (title, text, pages) == ("Bio Notes", "Lecture notes on ATP synthase.", None)


async def test_private_google_doc_explains_sharing():
    with respx.mock() as router:
        router.get("https://docs.google.com/document/d/P/export?format=txt").mock(return_value=httpx.Response(401))
        with pytest.raises(ingest.IngestError, match="Anyone with the link"):
            await ingest.fetch_web("https://docs.google.com/document/d/P/edit")


@needs_ocr
def test_scanned_pdf_is_ocrd():
    buf = io.BytesIO()
    text_image("Proton gradient drives ATP").save(buf, format="PDF")
    text, pages, ocr_count = ingest.extract_pdf(buf.getvalue())
    assert ocr_count == 1
    assert "Proton gradient" in text and "ATP" in text


@needs_ocr
async def test_image_source_uses_ocr_without_vision_model():
    buf = io.BytesIO()
    text_image("Krebs cycle").save(buf, format="PNG")
    meta: dict = {}
    text = await ingest.extract_image(buf.getvalue(), "slide.png", meta)
    assert "Krebs" in text and meta == {"ocr": True, "described": False}


def test_new_file_kinds():
    assert ingest.kind_for_filename("deck.PPTX") == "pptx"
    assert ingest.kind_for_filename("book.epub") == "epub"
    assert ingest.kind_for_filename("photo.jpeg") == "image"
