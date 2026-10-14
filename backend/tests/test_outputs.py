import io
import json
import zipfile

import httpx
import pytest
from fastapi.testclient import TestClient

from app import exporters, studio
from app.main import app
from tests.conftest import configure_llm
from tests.test_api import SOURCE_TEXT, signed_in, wait_for

DECK = {"title": "Cell Energy", "subtitle": "How cells make ATP", "slides": [
    {"layout": "section", "title": "Part 1"},
    {"layout": "bullets", "title": "Respiration", "bullets": ["Glycolysis", "Krebs cycle"], "notes": "Say hello"},
    {"layout": "two_column", "title": "Compare", "left": {"heading": "Aerobic", "bullets": ["O2"]}, "right": {"heading": "Anaerobic", "bullets": ["No O2"]}},
    {"layout": "quote", "quote": "Protons spin the turbine.", "attribution": "Lecture 7"},
    {"layout": "weird", "title": "Becomes bullets", "bullets": ["x"]},
    {"layout": "quote"},
]}


def test_slides_validate_and_pptx_roundtrip():
    from pptx import Presentation

    deck = studio.validate("slides", DECK)
    assert [s["layout"] for s in deck["slides"]] == ["section", "bullets", "two_column", "quote", "bullets"]
    prs = Presentation(io.BytesIO(exporters.build_pptx(deck)))
    assert len(prs.slides) == 1 + 5  # title slide + content
    texts = ["\n".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides]
    assert "Cell Energy" in texts[0] and "Respiration" in texts[2] and "Aerobic" in texts[3]
    assert prs.slides[2].notes_slide.notes_text_frame.text == "Say hello"


def test_infographic_validate_limits_counts():
    data = studio.validate("infographic", {"title": "T", "sections": [{"heading": str(i)} for i in range(6)],
                                           "stats": [{"value": "1"}, {"label": "no value"}], "steps": list("abcdefg")})
    assert len(data["sections"]) == 4 and data["stats"] == [{"value": "1"}] and len(data["steps"]) == 5
    with pytest.raises(ValueError):
        studio.validate("infographic", {"title": "T", "sections": []})


@pytest.fixture
def client(fake_llm):
    with TestClient(app) as c:
        yield c


def make_notebook(client, fake_llm):
    signed_in(client)
    configure_llm()
    fake_llm.replies = [("Summarise this source", "Summary")]
    nb = client.post("/api/notebooks", json={"title": "Bio"}).json()
    src = client.post(f"/api/notebooks/{nb['id']}/sources/text", json={"title": "Lecture 7", "text": SOURCE_TEXT}).json()
    wait_for(client, f"/api/sources/{src['id']}")
    return nb


def test_slides_infographic_share_and_export(client, fake_llm, monkeypatch, data_dir):
    nb = make_notebook(client, fake_llm)
    fake_llm.replies = [('"layout": "section"', json.dumps(DECK)),
                        ("At a glance", "unused"),
                        ('"takeaway"', json.dumps({"title": "Cell Energy", "sections": [{"heading": "ATP", "text": "Made by a turbine"}], "takeaway": "Gradients power cells"}))]

    async def fake_still(folder):
        (folder / "still.png").write_bytes(b"\x89PNG fake")
        assert json.loads((folder / "still.json").read_text())["kind"] == "infographic"
        return {"image": "still.png", "width": 1080, "height": 1920}

    monkeypatch.setattr("app.video.pipeline.render_still", fake_still)
    deck = wait_for(client, f"/api/artifacts/{client.post(f'/api/notebooks/{nb['id']}/artifacts', json={'kind': 'slides'}).json()['id']}")
    assert deck["status"] == "done" and deck["title"] == "Cell Energy", deck
    pptx = client.get(f"/api/artifacts/{deck['id']}/export/pptx")
    assert pptx.status_code == 200 and pptx.content[:2] == b"PK"
    assert "Cell-Energy.pptx" in pptx.headers["content-disposition"]

    info = wait_for(client, f"/api/artifacts/{client.post(f'/api/notebooks/{nb['id']}/artifacts', json={'kind': 'infographic'}).json()['id']}")
    assert info["status"] == "done", info
    assert info["data"]["image"].endswith("still.png")

    # Share links: public, read-only, scoped to the one output.
    token = client.post(f"/api/artifacts/{info['id']}/share").json()["token"]
    assert client.post(f"/api/artifacts/{info['id']}/share").json()["token"] == token
    anon = TestClient(app)
    shared = anon.get(f"/api/public/{token}").json()
    assert shared["kind"] == "infographic" and shared["notebook"] == "Bio"
    assert shared["data"]["image"] == f"/api/public/{token}/file/still.png"
    assert anon.get(shared["data"]["image"]).content == b"\x89PNG fake"
    assert anon.get(f"/api/public/{token}/file/../../sources").status_code == 404
    assert anon.get(f"/api/artifacts/{info['id']}").status_code == 401
    client.delete(f"/api/artifacts/{info['id']}/share")
    assert anon.get(f"/api/public/{token}").status_code == 404

    # Notebook export
    client.post(f"/api/notebooks/{nb['id']}/notes", json={"title": "My note", "content": "Remember the turbine"})
    resp = client.get(f"/api/notebooks/{nb['id']}/export")
    names = zipfile.ZipFile(io.BytesIO(resp.content)).namelist()
    assert "README.md" in names
    assert any(n.startswith("sources/01-Lecture-7") and n.endswith(".txt") for n in names)
    assert any(n.startswith("notes/") for n in names)
    assert any(n.endswith(".pptx") for n in names) and any(n.endswith(".png") for n in names)


def test_deep_research_reads_pages_and_recommends(client, fake_llm):
    nb = make_notebook(client, fake_llm)
    client.put("/api/settings", json={"search": {"base_url": "http://searx.test"}})
    page = "<html><head><title>{t}</title></head><body><article><p>" + "ATP synthase is a rotary motor. " * 20 + "</p></article></body></html>"
    fake_llm.replies = [
        ("Plan ", json.dumps({"queries": ["atp synthase mechanism", "chemiosmosis evidence"]})),
        ("Research question:", lambda body: json.dumps({"report": "ATP synthase rotates [1]. Evidence is strong [2].",
                                                        "recommended": [{"n": 1, "why": "clear explainer"}, {"n": 99}]})),
    ]
    results = {"results": [{"title": "Site A", "url": "http://site.test/a", "content": "a"},
                           {"title": "Site B", "url": "http://site.test/b", "content": "b"}]}
    fake_llm.router.get("http://searx.test/search").mock(return_value=httpx.Response(200, json=results))
    fake_llm.router.get("http://site.test/a").mock(return_value=httpx.Response(200, text=page.format(t="Page A"), headers={"content-type": "text/html"}))
    fake_llm.router.get("http://site.test/b").mock(return_value=httpx.Response(500))
    art = client.post(f"/api/notebooks/{nb['id']}/artifacts", json={"kind": "research", "options": {"query": "How does ATP synthase work?", "mode": "deep"}}).json()
    art = wait_for(client, f"/api/artifacts/{art['id']}", timeout=30)
    assert art["status"] == "done", art
    d = art["data"]
    assert d["mode"] == "deep" and len(d["queries"]) >= 2
    assert [r["read"] for r in d["references"]] == [True, False]
    assert d["recommended"] == [{"n": 1, "title": "Page A", "url": "http://site.test/a", "why": "clear explainer"}]
    assert "## References" in d["markdown"] and "[Page A](http://site.test/a)" in d["markdown"]


def test_research_without_search_explains(client, fake_llm):
    nb = make_notebook(client, fake_llm)
    fake_llm.replies = [("Plan ", json.dumps({"queries": ["q"]}))]
    art = client.post(f"/api/notebooks/{nb['id']}/artifacts", json={"kind": "research", "options": {"query": "Q?"}}).json()
    art = wait_for(client, f"/api/artifacts/{art['id']}")
    assert art["status"] == "failed" and "SearXNG" in art["error"]
