"""Text extraction for scanned pages, images, slides, e-books and Google Workspace links."""

from __future__ import annotations

import io
import os
import posixpath
import re
import shutil
import zipfile
from urllib.parse import urlparse
from xml.etree import ElementTree

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}
MIN_PAGE_CHARS = 25  # pages with less text than this are treated as scanned


class ExtractError(RuntimeError):
    pass


# ------------------------------------------------------------------ OCR

def ocr_available() -> bool:
    return shutil.which("tesseract") is not None


def ocr_image(image) -> str:
    """OCR a PIL image with Tesseract. Languages come from OCR_LANGS (e.g. "eng+spa")."""
    import pytesseract

    return pytesseract.image_to_string(image, lang=os.environ.get("OCR_LANGS", "eng")).strip()


def ocr_pdf_page(data: bytes, index: int, scale: float = 2.5) -> str:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(data)
    try:
        bitmap = pdf[index].render(scale=scale)
        return ocr_image(bitmap.to_pil())
    finally:
        pdf.close()


def ocr_image_bytes(data: bytes) -> str:
    from PIL import Image

    with Image.open(io.BytesIO(data)) as img:
        img.load()
        if img.mode not in ("RGB", "L"):
            img = img.convert("RGB")
        return ocr_image(img)


def image_mime(data: bytes, filename: str) -> str:
    from PIL import Image

    try:
        with Image.open(io.BytesIO(data)) as img:
            fmt = (img.format or "").lower()
    except Exception:  # noqa: BLE001
        fmt = ""
    fmt = fmt or filename.rsplit(".", 1)[-1].lower()
    return {"jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png", "webp": "image/webp",
            "gif": "image/gif"}.get(fmt, "image/png")


def to_png(data: bytes, max_side: int = 1600) -> bytes:
    """A PNG copy, downscaled, for sending to vision models (which reject TIFF/BMP and huge images)."""
    from PIL import Image

    with Image.open(io.BytesIO(data)) as img:
        img = img.convert("RGB")
        img.thumbnail((max_side, max_side))
        out = io.BytesIO()
        img.save(out, format="PNG")
        return out.getvalue()


# ------------------------------------------------------------------ PowerPoint

def extract_pptx(data: bytes) -> list[tuple[int, str]]:
    """Text of each slide (title, body, tables) plus its speaker notes."""
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    pages = []
    for number, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append("Speaker notes: " + notes)
        pages.append((number, "\n".join(parts) or "[slide with no text]"))
    if not any(not t.startswith("[slide with no text") for _, t in pages):
        raise ExtractError("This presentation has no text on its slides.")
    return pages


# ------------------------------------------------------------------ EPUB

def extract_epub(data: bytes) -> tuple[str, list[tuple[int, str]]]:
    """(title, chapters) following the book's reading order."""
    from lxml import html as lxml_html

    try:
        book = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise ExtractError("This EPUB file is damaged.") from exc
    ns = {"c": "urn:oasis:names:tc:opendocument:xmlns:container", "opf": "http://www.idpf.org/2007/opf",
          "dc": "http://purl.org/dc/elements/1.1/"}
    container = ElementTree.fromstring(book.read("META-INF/container.xml"))
    opf_path = container.find(".//c:rootfile", ns).attrib["full-path"]
    opf = ElementTree.fromstring(book.read(opf_path))
    base = posixpath.dirname(opf_path)
    title_el = opf.find(".//dc:title", ns)
    title = (title_el.text or "").strip() if title_el is not None else ""
    manifest = {item.attrib["id"]: item.attrib["href"] for item in opf.findall(".//opf:manifest/opf:item", ns)}
    chapters = []
    for itemref in opf.findall(".//opf:spine/opf:itemref", ns):
        href = manifest.get(itemref.attrib.get("idref", ""))
        if not href:
            continue
        try:
            raw = book.read(posixpath.normpath(posixpath.join(base, href)))
        except KeyError:
            continue
        doc = lxml_html.fromstring(raw)
        for bad in doc.xpath("//script|//style"):
            bad.getparent().remove(bad)
        blocks = [re.sub(r"\s+", " ", el.text_content()).strip()
                  for el in doc.xpath("//h1|//h2|//h3|//h4|//p|//li|//blockquote")]
        text = "\n\n".join(b for b in blocks if b)
        if text:
            chapters.append((len(chapters) + 1, text))
    if not chapters:
        raise ExtractError("No readable text found in this EPUB (it may be DRM-protected).")
    return title, chapters


# ------------------------------------------------------------------ Google Workspace links

GOOGLE_RE = re.compile(r"/(document|presentation|spreadsheets|file)/d/([\w-]+)")


def google_export_url(url: str) -> tuple[str, str] | None:
    """For a shared Google Docs/Slides/Sheets/Drive link: (download URL, kind of export)."""
    parsed = urlparse(url)
    if parsed.netloc not in ("docs.google.com", "drive.google.com"):
        return None
    match = GOOGLE_RE.search(parsed.path)
    if not match:
        return None
    kind, file_id = match.groups()
    if kind == "document":
        return f"https://docs.google.com/document/d/{file_id}/export?format=txt", "text"
    if kind == "presentation":
        return f"https://docs.google.com/presentation/d/{file_id}/export/txt", "text"
    if kind == "spreadsheets":
        return f"https://docs.google.com/spreadsheets/d/{file_id}/export?format=csv", "text"
    return f"https://drive.google.com/uc?export=download&id={file_id}", "file"


def filename_from_disposition(header: str) -> str | None:
    match = re.search(r"filename\*=UTF-8''([^;]+)", header) or re.search(r'filename="?([^";]+)"?', header)
    if not match:
        return None
    from urllib.parse import unquote

    return unquote(match.group(1)).strip()
