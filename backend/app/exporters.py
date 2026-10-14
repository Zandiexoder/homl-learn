"""Downloads: slide decks as PowerPoint, and whole notebooks as a zip."""

from __future__ import annotations

import io
import json
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from . import db
from .config import get_config

INK = (0x15, 0x20, 0x1B)
MUTED = (0x4A, 0x58, 0x52)
ACCENT = (0x0B, 0x6E, 0x55)
SOFT = (0xDD, 0xEF, 0xE7)
GROUND = (0xF7, 0xF9, 0xF6)


def build_pptx(deck: dict) -> bytes:
    """A 16:9 PowerPoint file from a generated deck (see studio 'slides')."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Emu, Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    W, H = prs.slide_width, prs.slide_height

    def rgb(c: tuple) -> RGBColor:
        return RGBColor(*c)

    def background(slide, color=GROUND) -> None:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = rgb(color)

    def box(slide, x, y, w, h, text: str, size: int, *, bold=False, color=INK, font="Calibri"):
        shape = slide.shapes.add_textbox(x, y, w, h)
        frame = shape.text_frame
        frame.word_wrap = True
        para = frame.paragraphs[0]
        para.text = text
        para.font.size, para.font.bold, para.font.name = Pt(size), bold, font
        para.font.color.rgb = rgb(color)
        return frame

    def bullets(slide, x, y, w, h, items: list, size: int = 22) -> None:
        frame = slide.shapes.add_textbox(x, y, w, h).text_frame
        frame.word_wrap = True
        for i, item in enumerate(items):
            para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            para.text = "•  " + str(item)
            para.font.size, para.font.name = Pt(size), "Calibri"
            para.font.color.rgb = rgb(INK)
            para.space_after = Pt(12)

    def accent_bar(slide) -> None:
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.55), Inches(1.4), Emu(70000))
        bar.fill.solid()
        bar.fill.fore_color.rgb = rgb(ACCENT)
        bar.line.fill.background()

    def notes(slide, text: str | None) -> None:
        if text:
            slide.notes_slide.notes_text_frame.text = str(text)

    # Title slide
    slide = prs.slides.add_slide(blank)
    background(slide, INK)
    box(slide, Inches(0.8), Inches(2.4), W - Inches(1.6), Inches(1.6), deck.get("title", "Slides"), 48, bold=True, color=(255, 255, 255))
    if deck.get("subtitle"):
        box(slide, Inches(0.8), Inches(4.1), W - Inches(1.6), Inches(1), deck["subtitle"], 22, color=SOFT)

    for s in deck.get("slides", []):
        slide = prs.slides.add_slide(blank)
        layout = s.get("layout", "bullets")
        if layout == "section":
            background(slide, ACCENT)
            box(slide, Inches(0.8), Inches(3), W - Inches(1.6), Inches(1.5), s.get("title", ""), 40, bold=True, color=(255, 255, 255))
        elif layout == "quote":
            background(slide)
            box(slide, Inches(1.2), Inches(1.6), W - Inches(2.4), Inches(3.4), "“" + str(s.get("quote", "")) + "”", 32, bold=True)
            if s.get("attribution"):
                box(slide, Inches(1.2), Inches(5.2), W - Inches(2.4), Inches(0.8), "— " + str(s["attribution"]), 18, color=MUTED)
        elif layout == "two_column":
            background(slide)
            box(slide, Inches(0.8), Inches(0.6), W - Inches(1.6), Inches(1), s.get("title", ""), 32, bold=True)
            accent_bar(slide)
            col_w = (W - Inches(2.2)) // 2
            for i, side in enumerate((s.get("left") or {}, s.get("right") or {})):
                x = Inches(0.8) + i * (col_w + Inches(0.6))
                box(slide, x, Inches(2), col_w, Inches(0.7), side.get("heading", ""), 24, bold=True, color=ACCENT)
                bullets(slide, x, Inches(2.8), col_w, Inches(4), side.get("bullets", []), 20)
        else:
            background(slide)
            box(slide, Inches(0.8), Inches(0.6), W - Inches(1.6), Inches(1), s.get("title", ""), 32, bold=True)
            accent_bar(slide)
            bullets(slide, Inches(0.8), Inches(2), W - Inches(1.6), Inches(5), s.get("bullets", []))
        notes(slide, s.get("notes"))

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


KNOWN_SUFFIXES = {".pdf", ".docx", ".pptx", ".epub", ".txt", ".md", ".png", ".jpg", ".jpeg", ".webp", ".gif",
                  ".mp3", ".wav", ".m4a", ".mp4", ".mov", ".webm", ".csv", ".html"}


def safe_name(text: str, fallback: str = "untitled") -> str:
    name = re.sub(r"[^\w\- ]+", "", text).strip().replace(" ", "-")[:60]
    return name or fallback


def notebook_zip(notebook_id: str) -> tuple[str, bytes]:
    """Everything in a notebook: sources (original files and text), notes, chat, and Studio outputs."""
    nb = db.fetch_one("SELECT * FROM notebooks WHERE id=?", (notebook_id,))
    sources = db.fetch_all("SELECT * FROM sources WHERE notebook_id=? ORDER BY created_at", (notebook_id,))
    notes = db.fetch_all("SELECT * FROM notes WHERE notebook_id=? ORDER BY created_at", (notebook_id,))
    messages = db.fetch_all("SELECT * FROM messages WHERE notebook_id=? ORDER BY created_at", (notebook_id,))
    artifacts = db.fetch_all("SELECT * FROM artifacts WHERE notebook_id=? AND status='done' ORDER BY created_at", (notebook_id,))
    files = get_config().files_dir
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        readme = [f"# {nb['title']}", "", f"Exported from Homl Learn on {datetime.now(timezone.utc):%Y-%m-%d %H:%M} UTC.", ""]
        if nb.get("summary"):
            readme += ["## Summary", "", nb["summary"], ""]
        readme += ["## Sources", ""] + [f"- {s['title']}" + (f" ({s['uri']})" if s.get("uri") else "") for s in sources]
        z.writestr("README.md", "\n".join(readme) + "\n")
        for i, s in enumerate(sources, start=1):
            title = Path(s["title"]).stem if Path(s["title"]).suffix.lower() in KNOWN_SUFFIXES else s["title"]
            stem = f"sources/{i:02d}-{safe_name(title)}"
            if s.get("text"):
                z.writestr(f"{stem}.txt", s["text"])
            path = s["meta"].get("path")
            if path and (files / notebook_id).resolve() in Path(path).resolve().parents:
                try:
                    z.write(path, f"{stem}-original{Path(path).suffix}")
                except OSError:
                    pass
        for i, n in enumerate(notes, start=1):
            z.writestr(f"notes/{i:02d}-{safe_name(n['title'])}.md", f"# {n['title']}\n\n{n['content']}\n")
        if messages:
            lines = []
            for m in messages:
                who = "You" if m["role"] == "user" else "Homl Learn"
                lines += [f"**{who}:** {m['content']}", ""]
                used = [c for c in m.get("citations") or [] if c.get("used")]
                for c in used:
                    lines.append(f"> [{c['n']}] {c['source_title']}: {c['text'][:300]}")
                if used:
                    lines.append("")
            z.writestr("chat.md", "\n".join(lines))
        for i, a in enumerate(artifacts, start=1):
            stem = f"studio/{i:02d}-{a['kind']}-{safe_name(a['title'])}"
            d = a["data"]
            if "markdown" in d:
                z.writestr(f"{stem}.md", d["markdown"])
            elif a["kind"] == "slides":
                z.writestr(f"{stem}.pptx", build_pptx(d))
            elif a["kind"] in ("audio", "video", "infographic"):
                rel = d.get("audio") or d.get("video") or d.get("image")
                if rel and (files / rel).is_file():
                    z.write(files / rel, f"{stem}{(files / rel).suffix}")
                if a["kind"] == "audio":
                    z.writestr(f"{stem}-transcript.md", "\n\n".join(f"**Host {t['speaker']}:** {t['text']}" for t in d.get("turns", [])))
            else:
                z.writestr(f"{stem}.json", json.dumps(d, indent=2, ensure_ascii=False))
    return f"{safe_name(nb['title'], 'notebook')}.zip", buf.getvalue()
